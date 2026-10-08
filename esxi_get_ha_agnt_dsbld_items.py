# Project Name  :  Virtual Infrastructure's Micro Reports  
# ------------------------------------------------------- 
# Module Name   :         esxi_get_ha_agnt_dsbld_items.py 
# Created by    :       									 Yuri P. Bodrov 
# Email         : 									 bodrovyp@hotmail.com 
# Phone Number  :         									 +79259929596 
# Telegram      :          										@YuriBodrov 
# LinkedIn Page : https://www.linkedin.com/in/yuribodrov/

"""
NOTE : This Module is designed to Collect ESXi Hypervisors with 
												    vSphere HA Agent in Disabled State.
"""
#########################################################################################
from sys import stdout             # Function for an Output Print Options
from dataclasses import dataclass  # Namespace for Operations with Data Classes
from time import sleep             # Function for the Time Delaying
from pyVmomi import vim            # Namespace for a Core Operations with VI Objects
import openpyxl                    # For an Operations with MS Excel File(s)

@dataclass # Main Class of this Module 'esxi_get_ha_agnt_dsbld_items.py'
class GetESXisHAAgentDisabledClass(): ###################################################
	pass
	connstate  : bool                # Is vCenter Server Connected : True/False
	vc_instance: vim.ServiceInstance # Passing SmartConnect Service Instance

	def get_esxi_ha_agent_dis_func(self): #################################################
		pass
		vi_content = self.vc_instance.RetrieveContent()
		if ((vi_content is not None) and (self.connstate)):
			container = vi_content.rootFolder # Root/Parent Directory Starting Point
			view_type = [vim.HostSystem]      # Looking for Object(s) > ESXi Host(s)
			recursive = True                  # Include Subfolders within Root/Parent Directory

			assert vi_content.viewManager is not None 
			container_view = vi_content.viewManager.CreateContainerView(container, view_type, \
																															 recursive)

			# Open and Initialize the Existing(!) MS Excel File
			xlsx_wb = openpyxl.load_workbook("report.xlsx")

			# Create a Thematic Sheet and Set it as Active
			if (not ("esxi_ha_agent_disabled" in xlsx_wb.sheetnames)):
				pass
				xlsx_wb.create_sheet(title="esxi_ha_agent_disabled")
				active_xlsx_sheet = xlsx_wb["esxi_ha_agent_disabled"] # type: ignore

				# Set the Column Headers of Active XLSX Sheet
				column_header_list = ["ESXi Server", "vSphere HA Agent State"]
				active_xlsx_sheet.append(column_header_list) 					# type: ignore
				xlsx_wb.save("report.xlsx")                  					# type: ignore
			else:
				pass
				active_xlsx_sheet = xlsx_wb["esxi_ha_agent_disabled"] # type: ignore
		
			stdout.write(f"[INFO] : HA Agent Disabled. Collecting ESXi Hostnames...")
			stdout.flush()
			for host in container_view.view:
				das_info = host.runtime.dasHostState									# type: ignore

				if (das_info and hasattr(das_info, "state")):
					pass
				else:
					pass
					xlsx_data = (host.name, "Disabled")
					active_xlsx_sheet.append(xlsx_data) 								# type: ignore
					xlsx_wb.save("report.xlsx")         								# type: ignore
			
			container_view.Destroy() # to Avoid Memory Accumulation in the vCenter Server
			print("Done.")
			print("")
			sleep(2)
