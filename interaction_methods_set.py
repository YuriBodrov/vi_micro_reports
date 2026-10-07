# Project Name  :  Virtual Infrastructure's Micro Reports  
# ------------------------------------------------------- 
# Module Name   :         		 interaction_methods_set.py 
# Created by    :       									 Yuri P. Bodrov 
# Email         : 									 bodrovyp@hotmail.com 
# Phone Number  :         									 +79259929596 
# Telegram      :          										@YuriBodrov 
# LinkedIn Page : https://www.linkedin.com/in/yuribodrov/

"""
NOTE : This Module is just a Set of Methods for facilitating Interaction between 
			 Components and Variables
"""
#########################################################################################
from dataclasses import dataclass # Namespace for Operations with Data Classes
from openpyxl import Workbook     # For an Operations with MS Excel File
from openpyxl.styles import Font	# This Class provides All Font Customization Types
from openpyxl import worksheet    # Test Import

@dataclass # Main Class of this Module 'interaction_methods_set.py'
class InteractMethodsClass(): ###########################################################
	pass
	tmp_workbook      : Workbook
	tmp_stylesheet    : worksheet.worksheet.Worksheet() # type: ignore
	tmp_workbook_name : str
	
	def add_text_style_func(self):
		pass
		# Create a New Sheet's Font Style
		new_font_style = Font(name = "Segoe UI", size = 10, bold = False, italic = False)

		# Iterate through all Filled Cells on the Worksheet
		for row in self.tmp_stylesheet.iter_rows\
			(min_row = 1, max_row = self.tmp_stylesheet.max_row, min_col = 1, \
			 max_col = self.tmp_stylesheet.max_column):
			pass
			for cell in row:
				cell.font = new_font_style
		self.tmp_workbook.save(self.tmp_workbook_name)