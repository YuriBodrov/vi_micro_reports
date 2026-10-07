# Project Name  :  Virtual Infrastructure's Micro Reports  
# ------------------------------------------------------- 
# Module Name   :       clu_get_vsphere_ha_dsbld_items.py 
# Created by    :       									 Yuri P. Bodrov 
# Email         : 									 bodrovyp@hotmail.com 
# Phone Number  :         									 +79259929596 
# Telegram      :          										@YuriBodrov 
# LinkedIn Page : https://www.linkedin.com/in/yuribodrov/

"""
NOTE : This Module is used for the VMware vSphere Clusters with 'vSphere HA
			 					  in Disabled State' Collection from VMware Infrastructure.
"""
#########################################################################################
from sys import stdout             # Function for an Output Print Options
from dataclasses import dataclass  # Namespace for Operations with Data Classes
from time import sleep             # Function for the Time Delaying
from pyVmomi import vim            # Namespace for a Core Operations with VI Objects
import openpyxl                    # For an Operations with MS Excel File(s)

@dataclass # Main Class of this Module 'clu_get_vsphere_ha_dsbld_items.py'
class GetClustersHADisabledClass(): #####################################################
	pass
	connstate  : bool                # Is vCenter Server Connected : True/False
	vc_instance: vim.ServiceInstance # Passing SmartConnect Service Instance

	def get_ha_disabled_clusters_func(self): ##############################################
		pass
		vi_content = self.vc_instance.RetrieveContent()
		if ((vi_content is not None) and (self.connstate)):
			container = vi_content.rootFolder        # Root/Parent Directory Starting Point
			view_type = [vim.ClusterComputeResource] # Looking for Object(s) > Cluster(s)
			recursive = True                         # Include SubDirs within Parent Directory

			assert vi_content.viewManager is not None 
			container_view = vi_content.viewManager.CreateContainerView(container, view_type, \
																															 recursive)

			# Open and Initialize the Existing(!) MS Excel File
			xlsx_wb = openpyxl.load_workbook("report.xlsx")

			# Create a Thematic Sheet and Set it as Active
			if (not ("clu_vsphere_ha_disabled" in xlsx_wb.sheetnames)):
				pass
				xlsx_wb.create_sheet(title="clu_vsphere_ha_disabled")
				active_xlsx_sheet = xlsx_wb["clu_vsphere_ha_disabled"] 		# type: ignore

				# Set the Column Headers of Active XLSX Sheet
				column_header_list = ["Cluster Name", "vSphere HA Status"]
				active_xlsx_sheet.append(column_header_list) 							# type: ignore
				xlsx_wb.save("report.xlsx")                  							# type: ignore
			else:
				pass
				active_xlsx_sheet = xlsx_wb["clu_vsphere_ha_disabled"] 		# type: ignore
					
			stdout.write(f"[INFO] : HA Disabled. Collecting Cluster Names...")
			stdout.flush()

			for cluster in container_view.view:
				# Check If HA Enabled in the Particular Cluster
				ha_enabled_bool = cluster.configuration.dasConfig.enabled # type: ignore
				if (not ha_enabled_bool):
					xlsx_data = (cluster.name, "Disabled")
					active_xlsx_sheet.append(xlsx_data) 										# type: ignore
					xlsx_wb.save("report.xlsx")         										# type: ignore
				 
			container_view.Destroy() # to Avoid Memory Accumulation in the vCenter Server
			print("Done.")
			sleep(2)
			print("")
