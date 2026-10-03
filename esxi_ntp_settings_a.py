# Project Name  : 										Report-as-a-Service  
# ------------------------------------------------------- 
# Module Name   :         					 esxi_ntp_settings.py 
# Created by    :       									 Yuri P. Bodrov 
# Email         : 									 bodrovyp@hotmail.com 
# Phone Number  :         									 +79259929596 
# Telegram      :          										@YuriBodrov 
# LinkedIn Page : https://www.linkedin.com/in/yuribodrov/

"""
NOTE : This module is designed to Detect NTP Service Configuration Deviations Within
			 																													the ESXi Hypervisors.
"""
#########################################################################################
from sys import stdout            # Function for an Output Print Options
from dataclasses import dataclass # Function for Operations with Data Classes
from time import sleep            # Function for the Time Delaying "sleep()" Function
from pyVmomi import vim           # Namespace for a Core Operations with VI Objects
import openpyxl                   # For an Operations with MS Excel File(s)

@dataclass # Main Class of this Module 'esxi_ntp_settings.py'
class GetEsxiNtpSettings(): #############################################################
	pass
	connstate  : bool                # Is vCenter Server Connected : True/False
	vc_instance: vim.ServiceInstance # Passing SmartConnect Service Instance
	esxi_ntp_set_miss = []

	def get_esxi_ntp_set_func(self): ######################################################
		pass
		vi_content = self.vc_instance.RetrieveContent()
		if ((vi_content is not None) and (self.connstate)):
			container = vi_content.rootFolder # Root/Parent Directory Starting Point
			view_type = [vim.HostSystem]      # Look for Object(s) : Host(s)
			recursive = True                  # Include SubDirs within Root/Parent Directory

			"""
			NOTE : Check this Out! 'Assert' Keyword. We are preventing the return of a 
			      'Null/None' Function's Value here
			"""
			assert vi_content.viewManager is not None 
			container_view = vi_content.viewManager.CreateContainerView(container, \
												view_type, recursive)
		
			# Open and Initialize the Existing(!) MS Excel File
			xlsx_wb = openpyxl.load_workbook("report.xlsx")

			# Create a Thematic Sheet and Set it as Active
			if (not ("esxi_ntp_anomalies" in xlsx_wb.sheetnames)):
				xlsx_wb.create_sheet(title="esxi_ntp_anomalies")
				active_xlsx_sheet = xlsx_wb["esxi_ntp_anomalies"] # type: ignore

				# Set the Column Headers of Active XLSX Sheet
				column_header_list = ["ESXi Server", "Service Name", "Service Config", \
													"Service State"]
				active_xlsx_sheet.append(column_header_list) # type: ignore
				xlsx_wb.save("report.xlsx")                  # type: ignore
			else:
				active_xlsx_sheet = xlsx_wb["esxi_ntp_anomalies"] # type: ignore

			stdout.write(f"[INFO] : Collecting ESXi Server's NTP Settings...")
			stdout.flush()
			
			for host in container_view.view:
				# Fetch NTP Servers and Time Configuration
				# Ref : vim.host.config.dateTimeInfo
				"""
				NOTE : '[host_system].config.dateTimeInfo.ntpConfig.server' : This retrieves
							 a List of Strings representing the NTP IP addresses or FQDNs mapped to the
							 ESXi Server
				"""
				esxi_datetime_info = host.config.dateTimeInfo
				if ((esxi_datetime_info) and (esxi_datetime_info.ntpConfig)):
					ntp_servers = esxi_datetime_info.ntpConfig.server
				else:
					#print(f"{host.name} : No NTP Servers so far!")
					pass
					xlsx_data = (host.name, "ntpd", "unknown", "unknown")
					active_xlsx_sheet.append(xlsx_data) # type: ignore
					xlsx_wb.save("report.xlsx")         # type: ignore

				# Fetch NTP Service Status and Startup Policy
				# Reference to 'vim.host.configManager.serviceSystem'
				"""
				NOTE : '[ntp_service].policy' Returns the Startup Policy. 
							 Common values include :
							 --- 
							 'on'        : Start and Stop with Host, 
							 'off'       : Manual startup, 
							 'automatic' : Start automatically if network ports are open. 
							 ---
							 '[ntp_service].running' : a Boolean Value (True/False) telling us whether
																				 the 'ntpd' Service is actively running now. 
							 '[ntp_service].key'     : Service Name. 
				"""
				esxi_services_data = host.configManager.serviceSystem
				if ((esxi_services_data) and (esxi_services_data.serviceInfo)):
					for s_name in esxi_services_data.serviceInfo.service:
						if (s_name.key == "ntpd"): # If Service Name is 'ntpd'
							if (not s_name.running): 
								#print(f"{host.name} | {s_name.key} | {s_name.policy} | {s_name.running}")
								if (s_name.running == False):
									s_name_run_state = "Stopped"
								xlsx_data = (host.name, s_name.key, s_name.policy, s_name_run_state)
								active_xlsx_sheet.append(xlsx_data) # type: ignore
								xlsx_wb.save("report.xlsx")         # type: ignore
				 
			print("Done.")
			sleep(2)
			print("")

